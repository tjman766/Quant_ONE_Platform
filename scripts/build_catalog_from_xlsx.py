"""키움 REST API 문서(xlsx)를 api_catalog.json으로 변환하는 보조 스크립트.
토큰 API(au10001/au10002)는 문서 보존용으로만 남기고 자동 화면/거래 코딩 대상에서는 제외합니다.
"""
from __future__ import annotations
import json, re, sys
from pathlib import Path
from openpyxl import load_workbook

TOKEN_IDS = {"au10001", "au10002"}

def norm(v):
    return "" if v is None else str(v).strip()

def find_api_id(text: str) -> str:
    m = re.search(r"\(([A-Za-z0-9]+)\)", text or "")
    return m.group(1) if m else ""

def build(xlsx: Path) -> dict:
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    apis=[]
    for ws in wb.worksheets:
        name=ws.title
        api_id=find_api_id(name)
        if not api_id: continue
        values=[[norm(c.value) for c in row] for row in ws.iter_rows(max_row=160, max_col=8)]
        flat="\n".join(" ".join(r) for r in values)
        kind="realtime" if ("WebSocket" in flat or "실시간" in name or "REG" in flat) else "rest"
        method="POST" if "POST" in flat.upper() else ""
        url=""
        for r in values:
            for c in r:
                if c.startswith("/api/") or c.startswith("/oauth"):
                    url=c
                    break
            if url: break
        apis.append({"sheet_name":name,"api_id":api_id,"name":re.sub(r"\([^)]+\)","",name).strip(),"kind":kind,"method":method,"url":url,"auto_code_target":api_id not in TOKEN_IDS})
    return {"source_file":xlsx.name,"summary":{"total":len(apis),"realtime":sum(a['kind']=='realtime' for a in apis),"rest":sum(a['kind']=='rest' for a in apis),"token_excluded":len([a for a in apis if a['api_id'] in TOKEN_IDS])},"apis":apis}

if __name__ == "__main__":
    xlsx=Path(sys.argv[1])
    out=Path(sys.argv[2]) if len(sys.argv)>2 else Path("data/kiwoom_api_catalog.json")
    out.write_text(json.dumps(build(xlsx),ensure_ascii=False,indent=2),encoding="utf-8")
    print(out)
